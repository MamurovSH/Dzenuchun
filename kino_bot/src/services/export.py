"""
Export service (#11) — CSV, Excel (xlsx), JSON exports for admin.
"""
from __future__ import annotations

import csv
import io
import json
from datetime import datetime
from typing import List

from src.db.models.movie import Movie
from src.db.models.user import User


class ExportService:

    # ── Movies ────────────────────────────────────────────────────────────

    @staticmethod
    def movies_to_csv(movies: List[Movie]) -> bytes:
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow([
            "ID", "Code", "Title", "Title_UZ", "Title_RU", "Title_EN",
            "Year", "Genres", "Director", "Country", "Language",
            "Rating", "Duration", "View Count", "Is Active", "Created At",
        ])
        for m in movies:
            writer.writerow([
                m.id, m.code, m.title, m.title_uz or "", m.title_ru or "", m.title_en or "",
                m.year, ", ".join(m.genres or []), m.director or "",
                m.country or "", m.language or "",
                m.rating, m.duration or "",
                m.view_count, m.is_active,
                str(m.created_at),
            ])
        return buf.getvalue().encode("utf-8-sig")

    @staticmethod
    def movies_to_json(movies: List[Movie]) -> bytes:
        data = []
        for m in movies:
            data.append({
                "id": m.id, "code": m.code,
                "title": m.title, "title_uz": m.title_uz, "title_ru": m.title_ru,
                "title_en": m.title_en, "year": m.year,
                "genres": m.genres or [], "director": m.director,
                "country": m.country, "language": m.language,
                "rating": m.rating, "duration": m.duration,
                "view_count": m.view_count, "is_active": m.is_active,
                "created_at": str(m.created_at),
            })
        return json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")

    @staticmethod
    def movies_to_excel(movies: List[Movie]) -> bytes:
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            raise RuntimeError("openpyxl required: pip install openpyxl")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Movies"

        headers = [
            "ID", "Code", "Title", "Uzbek Title", "Russian Title", "English Title",
            "Year", "Genres", "Director", "Country", "Language",
            "Rating", "Duration", "Views", "Active", "Created At",
        ]
        header_fill = PatternFill("solid", fgColor="1F4E79")
        header_font = Font(color="FFFFFF", bold=True)

        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for row, m in enumerate(movies, 2):
            ws.append([
                m.id, m.code, m.title, m.title_uz or "", m.title_ru or "", m.title_en or "",
                m.year, ", ".join(m.genres or []), m.director or "",
                m.country or "", m.language or "",
                m.rating, m.duration or "",
                m.view_count, m.is_active, str(m.created_at),
            ])

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # ── Users ─────────────────────────────────────────────────────────────

    @staticmethod
    def users_to_csv(users: List[User]) -> bytes:
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow([
            "ID", "Telegram ID", "Username", "Full Name", "Language", "Status",
            "Total Searches", "Total Views", "Warn Count", "Created At", "Last Active",
        ])
        for u in users:
            writer.writerow([
                u.id, u.telegram_id, u.username or "", u.full_name,
                u.language.value, u.status.value,
                u.total_searches, u.total_views, u.warn_count,
                str(u.created_at), str(u.last_active or ""),
            ])
        return buf.getvalue().encode("utf-8-sig")

    @staticmethod
    def users_to_excel(users: List[User]) -> bytes:
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            raise RuntimeError("openpyxl required: pip install openpyxl")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Users"
        headers = [
            "ID", "Telegram ID", "Username", "Full Name", "Language", "Status",
            "Searches", "Views", "Warns", "Created At", "Last Active",
        ]
        header_fill = PatternFill("solid", fgColor="1F4E79")
        header_font = Font(color="FFFFFF", bold=True)
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill = header_fill
            cell.font = header_font

        for u in users:
            ws.append([
                u.id, u.telegram_id, u.username or "", u.full_name,
                u.language.value, u.status.value,
                u.total_searches, u.total_views, u.warn_count,
                str(u.created_at), str(u.last_active or ""),
            ])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # ── Statistics ────────────────────────────────────────────────────────

    @staticmethod
    def stats_to_json(stats: dict) -> bytes:
        stats["exported_at"] = datetime.utcnow().isoformat()
        return json.dumps(stats, ensure_ascii=False, indent=2).encode("utf-8")
